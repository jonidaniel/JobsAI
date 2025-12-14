# ---------- TESTS FOR GENERATOR AGENT ----------

import os
import pytest
from unittest.mock import patch, MagicMock
from datetime import datetime
from docx import Document

from jobsai.agents.generator import GeneratorAgent

# Mock data
mock_profile = "Experienced Python developer with 5 years of experience in web development and AI/ML."

mock_job_analysis_single = """
Job Analysis
========================================
Top 1 Jobs:

Title: Senior Python Developer
Company: Company A
Location: Helsinki
Score: 85%
Matched Skills: Python, JavaScript
Missing Skills: Docker, Kubernetes
URL: https://example.com/job1
Instructions: Focus on Python experience and web development background.
----------------------------------------
"""

mock_job_analysis_multiple = """
Job Analysis
========================================
Top 2 Jobs:

Title: Senior Python Developer
Company: Company A
Location: Helsinki
Score: 85%
Matched Skills: Python, JavaScript
Missing Skills: Docker, Kubernetes
URL: https://example.com/job1
Instructions: Focus on Python experience.
----------------------------------------
Title: AI Engineer
Company: Company B
Location: Espoo
Score: 75%
Matched Skills: Python, AI/ML
Missing Skills: TensorFlow
URL: https://example.com/job2
Instructions: Emphasize AI/ML background.
----------------------------------------
"""

mock_llm_cover_letter = """
Dear Hiring Manager,

I am writing to express my interest in the Senior Python Developer position at Company A.
With 5 years of experience in Python and web development, I am excited about this opportunity.

I have extensive experience with Python and JavaScript, which aligns well with your requirements.
I am eager to learn Docker and Kubernetes to further enhance my skills.

Thank you for considering my application.

Best regards,
[Your Name]
"""


@pytest.fixture
def generator():
    """Create a GeneratorAgent instance for testing."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return GeneratorAgent(timestamp)


@pytest.fixture(autouse=True)
def clean_cover_letters_folder():
    """Clean up cover letters folder before and after tests."""
    from jobsai.config.paths import COVER_LETTER_PATH

    folder = COVER_LETTER_PATH
    if os.path.exists(folder):
        for f in os.listdir(folder):
            file_path = os.path.join(folder, f)
            if os.path.isfile(file_path):
                os.remove(file_path)
    else:
        os.makedirs(folder, exist_ok=True)
    yield
    # Cleanup after test
    if os.path.exists(folder):
        for f in os.listdir(folder):
            file_path = os.path.join(folder, f)
            if os.path.isfile(file_path):
                os.remove(file_path)


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_basic(mock_call_llm, generator):
    """Test that cover letters are generated."""
    letters = generator.generate_letters(
        mock_job_analysis_single, mock_profile, "Professional", num_letters=1
    )
    assert isinstance(letters, list)
    assert len(letters) == 1
    assert isinstance(letters[0], Document)
    assert mock_call_llm.called


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_multiple(mock_call_llm, generator):
    """Test generating multiple cover letters."""
    letters = generator.generate_letters(
        mock_job_analysis_multiple, mock_profile, "Professional", num_letters=2
    )
    assert len(letters) == 2
    assert all(isinstance(letter, Document) for letter in letters)
    # Should call LLM once per letter
    assert mock_call_llm.call_count == 2


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_limits_to_num_letters(mock_call_llm, generator):
    """Test that generation is limited to num_letters."""
    letters = generator.generate_letters(
        mock_job_analysis_multiple, mock_profile, "Professional", num_letters=1
    )
    assert len(letters) == 1
    # Should only call LLM once
    assert mock_call_llm.call_count == 1


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_saves_to_disk(mock_call_llm, generator):
    """Test that cover letters are saved to disk."""
    from jobsai.config.paths import COVER_LETTER_PATH

    letters = generator.generate_letters(
        mock_job_analysis_single, mock_profile, "Professional", num_letters=1
    )
    # Check that file was created
    files = os.listdir(COVER_LETTER_PATH)
    assert len(files) > 0
    # Find files with our timestamp
    matching_files = [f for f in files if generator.timestamp in f]
    assert len(matching_files) > 0


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_document_structure(mock_call_llm, generator):
    """Test that generated documents have correct structure."""
    letters = generator.generate_letters(
        mock_job_analysis_single, mock_profile, "Professional", num_letters=1
    )
    doc = letters[0]
    # Should have paragraphs (contact info, date, recipient, body, signature)
    assert len(doc.paragraphs) > 0
    # Check for placeholders
    text = "\n".join([p.text for p in doc.paragraphs])
    assert "ADD" in text  # Placeholders like "ADD EMAIL", "ADD RECRUITER", etc.


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_handles_style_string(mock_call_llm, generator):
    """Test that style can be a string."""
    letters = generator.generate_letters(
        mock_job_analysis_single, mock_profile, "Friendly", num_letters=1
    )
    assert len(letters) == 1
    assert mock_call_llm.called


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_handles_style_list(mock_call_llm, generator):
    """Test that style can be a list (uses first style)."""
    letters = generator.generate_letters(
        mock_job_analysis_single,
        mock_profile,
        ["Professional", "Friendly"],
        num_letters=1,
    )
    assert len(letters) == 1
    assert mock_call_llm.called


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_defaults_to_professional(mock_call_llm, generator):
    """Test that unknown style defaults to Professional."""
    letters = generator.generate_letters(
        mock_job_analysis_single, mock_profile, "UnknownStyle", num_letters=1
    )
    assert len(letters) == 1
    # Should still work with default style
    assert mock_call_llm.called


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_parses_job_sections(mock_call_llm, generator):
    """Test that job analysis is parsed into sections."""
    letters = generator.generate_letters(
        mock_job_analysis_multiple, mock_profile, "Professional", num_letters=2
    )
    # Should generate 2 letters from 2 job sections
    assert len(letters) == 2
    assert mock_call_llm.call_count == 2


@patch("jobsai.utils.llms.call_llm", return_value=mock_llm_cover_letter)
def test_generate_letters_handles_empty_style_list(mock_call_llm, generator):
    """Test that empty style list defaults to Professional."""
    letters = generator.generate_letters(
        mock_job_analysis_single, mock_profile, [], num_letters=1
    )
    assert len(letters) == 1
    assert mock_call_llm.called
